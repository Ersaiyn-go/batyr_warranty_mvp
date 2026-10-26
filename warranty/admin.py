import csv
from datetime import datetime
from decimal import Decimal

from django import forms
from django.contrib import admin, messages
from django.db import transaction
from django.db.models import Count
from django.shortcuts import redirect, render
from django.urls import path

from openpyxl import load_workbook

from .models import ProductWarranty


class DuplicateSerialFilter(admin.SimpleListFilter):
    title = 'Дубли серийника'
    parameter_name = 'duplicate'

    def lookups(self, request, model_admin):
        return [('yes', 'Только дубли'), ('no', 'Без дублей')]

    def queryset(self, request, queryset):
        duplicated = (
            ProductWarranty.objects.values('serial_number')
            .annotate(cnt=Count('id'))
            .filter(cnt__gt=1)
            .values('serial_number')
        )
        if self.value() == 'yes':
            return queryset.filter(serial_number__in=duplicated)
        if self.value() == 'no':
            return queryset.exclude(serial_number__in=duplicated)
        return queryset


class DuplicateImportForm(forms.Form):
    file = forms.FileField(
        label='Файл с дублями',
        help_text='.txt / .tsv (скопировано из таблицы), .csv или .xlsx'
    )
    dry_run = forms.BooleanField(
        label='Пробный запуск (только показать, ничего не сохранять)',
        required=False,
        initial=True,
    )


class WarrantyImportForm(forms.Form):
    file = forms.FileField(
        label='Файл Excel или CSV',
        help_text='Поддерживаются файлы .xlsx и .csv'
    )


def clean_value(value):
    if value is None:
        return ''
    return str(value).strip()


def normalize_header(value):
    return clean_value(value).lower().replace('ё', 'е')


def parse_date(value):
    if value in [None, '']:
        return None

    if hasattr(value, 'date'):
        return value.date()

    value = clean_value(value)

    formats = [
        '%d.%m.%Y',
        '%Y-%m-%d',
        '%d/%m/%Y',
        '%d-%m-%Y',
    ]

    for fmt in formats:
        try:
            return datetime.strptime(value, fmt).date()
        except ValueError:
            pass

    return None


def parse_int(value, default=12):
    if value in [None, '']:
        return default

    try:
        return int(Decimal(str(value).strip()))
    except Exception:
        return default


def normalize_status(value, sale_date=None, buyer_name='', buyer_phone=''):
    value = clean_value(value).lower()

    status_map = {
        'sold': ProductWarranty.Status.SOLD,
        'продан': ProductWarranty.Status.SOLD,
        'продано': ProductWarranty.Status.SOLD,

        'in_stock': ProductWarranty.Status.IN_STOCK,
        'на складе': ProductWarranty.Status.IN_STOCK,
        'склад': ProductWarranty.Status.IN_STOCK,

        'returned': ProductWarranty.Status.RETURNED,
        'возврат': ProductWarranty.Status.RETURNED,

        'replaced': ProductWarranty.Status.REPLACED,
        'замена': ProductWarranty.Status.REPLACED,
        'заменен': ProductWarranty.Status.REPLACED,

        'blocked': ProductWarranty.Status.BLOCKED,
        'заблокирован': ProductWarranty.Status.BLOCKED,
    }

    if value in status_map:
        return status_map[value]

    if sale_date or buyer_name or buyer_phone:
        return ProductWarranty.Status.SOLD

    return ProductWarranty.Status.IN_STOCK


def get_column(row_dict, *names):
    for name in names:
        key = normalize_header(name)
        if key in row_dict:
            return row_dict[key]
    return ''

def normalize_product_name(value):
    return clean_value(value).lower().replace('ё', 'е').strip()


def get_auto_color(model_name):
    name = normalize_product_name(model_name)

    titanium_models = {
        'сункар про',
        'сұңқар про',
        'sunqar pro',
    }

    black_models = {
        'тенгер',
        'tenger',
        'қыран',
        'кыран',
        'kyran',
        'batyr pro',
        'батыр про',
        'sunqar black',
        'sunqar health',
        'сұңқар блек',
        'сункар блек',
        'health',
    }

    gray_models = {
        'томирис',
        'tomiris',
    }

    beige_models = {
        'томи',
        'tomi',
    }

    if name in titanium_models:
        return 'Титан'

    if name in black_models:
        return 'Черный'

    if name in gray_models:
        return 'Серый'

    if name in beige_models:
        return 'Бежевый'

    return ''


def get_auto_warranty_months(model_name):
    name = normalize_product_name(model_name)

    six_month_models = {
        'sunqar black',
        'sunqar health',
        'сұңқар блек',
        'сункар блек',
        'health',
        'томи',
        'tomi',
    }

    if name in six_month_models:
        return 6

    return 12

def read_xlsx(file):
    workbook = load_workbook(file, read_only=True, data_only=True)
    sheet = workbook.active

    rows = list(sheet.iter_rows(values_only=True))

    if not rows:
        return []

    headers = [normalize_header(cell) for cell in rows[0]]
    result = []

    for row in rows[1:]:
        row_dict = {}

        for index, header in enumerate(headers):
            if not header:
                continue

            value = row[index] if index < len(row) else ''
            row_dict[header] = value

        result.append(row_dict)

    return result


def read_csv(file):
    decoded_file = file.read().decode('utf-8-sig').splitlines()

    sample = '\n'.join(decoded_file[:5])
    delimiter = ';' if sample.count(';') >= sample.count(',') else ','

    reader = csv.DictReader(decoded_file, delimiter=delimiter)
    result = []

    for row in reader:
        normalized = {}

        for key, value in row.items():
            normalized[normalize_header(key)] = value

        result.append(normalized)

    return result


@admin.register(ProductWarranty)
class ProductWarrantyAdmin(admin.ModelAdmin):
    change_list_template = 'admin/warranty/productwarranty/change_list.html'

    list_display = (
        'serial_number',
        'model_name',
        'buyer_name',
        'buyer_phone',
        'sale_date',
        'warranty_months',
        'status',
    )

    search_fields = (
        'serial_number',
        'model_name',
        'buyer_name',
        'buyer_phone',
    )

    list_filter = (
        DuplicateSerialFilter,
        'status',
        'model_name',
        'sale_date',
    )

    ordering = ('-created_at',)

    def get_urls(self):
        urls = super().get_urls()

        custom_urls = [
            path(
                'import/',
                self.admin_site.admin_view(self.import_file),
                name='warranty_productwarranty_import',
            ),
            path(
                'import-duplicates/',
                self.admin_site.admin_view(self.import_duplicates),
                name='warranty_productwarranty_import_duplicates',
            ),
        ]

        return custom_urls + urls

    def import_file(self, request):
        if request.method == 'POST':
            form = WarrantyImportForm(request.POST, request.FILES)

            if form.is_valid():
                uploaded_file = form.cleaned_data['file']
                filename = uploaded_file.name.lower()

                try:
                    if filename.endswith('.xlsx'):
                        rows = read_xlsx(uploaded_file)
                    elif filename.endswith('.csv'):
                        rows = read_csv(uploaded_file)
                    else:
                        messages.error(request, 'Загрузите файл в формате .xlsx или .csv')
                        return redirect('..')

                    created_count = 0
                    updated_count = 0
                    skipped_count = 0
                    duplicate_skipped = []
                    errors = []

                    with transaction.atomic():
                        for index, row in enumerate(rows, start=2):
                            serial_number = clean_value(
                                get_column(
                                    row,
                                    'серийный номер',
                                    'serial_number',
                                    'serial',
                                    'серийник'
                                )
                            ).upper()

                            if not serial_number:
                                skipped_count += 1
                                continue

                            buyer_name = clean_value(
                                get_column(
                                    row,
                                    'фио владельца',
                                    'фио',
                                    'покупатель',
                                    'buyer_name',
                                    'owner_name'
                                )
                            )

                            buyer_phone = clean_value(
                                get_column(
                                    row,
                                    'номер владельца',
                                    'телефон',
                                    'номер телефона',
                                    'buyer_phone',
                                    'owner_phone'
                                )
                            )

                            model_name = clean_value(
                                get_column(
                                    row,
                                    'имя часов',
                                    'модель',
                                    'model_name',
                                    'watch_name'
                                )
                            )

                            excel_color = clean_value(
                                get_column(
                                    row,
                                    'цвет',
                                    'color'
                                )
                            )

                            auto_color = get_auto_color(model_name)

                            if auto_color:
                                color = auto_color
                            else:
                                color = excel_color

                            sale_date = parse_date(
                                get_column(
                                    row,
                                    'дата покупки',
                                    'дата продажи',
                                    'sale_date',
                                    'purchase_date'
                                )
                            )

                            warranty_months = get_auto_warranty_months(model_name)

                            status = normalize_status(
                                get_column(row, 'статус', 'status'),
                                sale_date=sale_date,
                                buyer_name=buyer_name,
                                buyer_phone=buyer_phone,
                            )

                            if not model_name:
                                model_name = 'BR Watch'

                            # Серийник, выданный двум покупателям, обычным импортом не трогаем,
                            # иначе непонятно, чью запись перезаписывать.
                            if ProductWarranty.objects.filter(
                                serial_number__iexact=serial_number
                            ).count() > 1:
                                duplicate_skipped.append(serial_number)
                                continue

                            product, created = ProductWarranty.objects.update_or_create(
                                serial_number=serial_number,
                                defaults={
                                    'buyer_name': buyer_name,
                                    'buyer_phone': buyer_phone,
                                    'model_name': model_name,
                                    'color': color,
                                    'sale_date': sale_date,
                                    'warranty_months': warranty_months,
                                    'status': status,
                                }
                            )

                            if created:
                                created_count += 1
                            else:
                                updated_count += 1

                    messages.success(
                        request,
                        f'Импорт завершён. Добавлено: {created_count}. '
                        f'Обновлено: {updated_count}. Пропущено: {skipped_count}.'
                    )

                    if duplicate_skipped:
                        messages.warning(
                            request,
                            'Не изменены серийники с двумя покупателями (правьте их вручную '
                            'или через «Импорт дублей»): ' + ', '.join(duplicate_skipped)
                        )

                    return redirect('..')

                except Exception as error:
                    messages.error(request, f'Ошибка импорта: {error}')
                    return redirect('..')

        else:
            form = WarrantyImportForm()

        context = {
            **self.admin_site.each_context(request),
            'title': 'Импорт серийных номеров',
            'form': form,
        }

        return render(
            request,
            'admin/warranty/productwarranty/import.html',
            context
        )

    def import_duplicates(self, request):
        from .duplicates import apply_duplicates, parse_duplicate_rows, read_rows

        result = None

        if request.method == 'POST':
            form = DuplicateImportForm(request.POST, request.FILES)

            if form.is_valid():
                uploaded_file = form.cleaned_data['file']
                dry_run = form.cleaned_data['dry_run']

                try:
                    rows = read_rows(uploaded_file, uploaded_file.name)
                    items, problems = parse_duplicate_rows(rows)
                    stats, log = apply_duplicates(items, dry_run=dry_run)
                    result = {
                        'dry_run': dry_run,
                        'stats': stats,
                        'log': log,
                        'problems': problems,
                    }
                    if not dry_run:
                        messages.success(
                            request,
                            f'Готово. Серийников: {stats["serials"]}. '
                            f'Обновлено: {stats["updated"]}. Создано: {stats["created"]}.'
                        )
                except Exception as error:
                    messages.error(request, f'Ошибка импорта: {error}')
        else:
            form = DuplicateImportForm()

        context = {
            **self.admin_site.each_context(request),
            'title': 'Импорт дублей серийных номеров',
            'form': form,
            'result': result,
        }

        return render(
            request,
            'admin/warranty/productwarranty/import_duplicates.html',
            context
        )
