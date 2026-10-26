"""
Загрузка серийных номеров, которые по ошибке выдали двум покупателям.

Формат файла (как в таблице, по колонкам, без заголовка или с заголовком):
  1  Серийный номер
  2  Дата покупки (покупатель 1)
  3  Модель (покупатель 1)
  4  Имя (покупатель 1)
  5  Телефон (покупатель 1)
  6  (пусто — разделитель)
  7  Дата покупки (покупатель 2)
  8  Модель (покупатель 2)
  9  Имя (покупатель 2)
  10 Телефон (покупатель 2)

Поддерживаются .txt/.tsv (через Tab, как копируется из Excel/Google Sheets),
.csv (через ; или ,) и .xlsx.

Логика для каждого серийника: если в базе уже есть запись с этим именем
(или телефоном) — она обновляется; иначе обновляется свободная запись
с этим серийником; если свободных нет — создаётся новая. Ничего не удаляется.
"""
import csv
import io
import re

from django.db import transaction

from .models import ProductWarranty


def _helpers():
    # Импорт внутри функции, чтобы не было циклического импорта с admin.py
    from .admin import get_auto_color, get_auto_warranty_months, parse_date
    return get_auto_color, get_auto_warranty_months, parse_date

SERIAL_RE = re.compile(r'^[A-Z]{2,}[A-Z0-9-]*\d+$')
INVISIBLE = dict.fromkeys(map(ord, '‪‫‬‭‮‎‏​﻿'), None)

DUPLICATE_NOTE = 'Серийный номер по ошибке выдан двум покупателям — на сайте показывается выбор покупки.'


def _clean(value):
    if value is None:
        return ''
    if hasattr(value, 'strftime'):
        return value
    text = str(value).translate(INVISIBLE)
    text = ' '.join(text.split())
    return text.strip()


def _digits(value):
    return ''.join(ch for ch in str(value or '') if ch.isdigit())


def _norm_name(value):
    return ' '.join(str(value or '').lower().replace('ё', 'е').replace('.', ' ').split())


def read_rows(uploaded_file, filename):
    """Возвращает список строк (списков ячеек)."""
    filename = filename.lower()

    if filename.endswith('.xlsx'):
        from openpyxl import load_workbook
        workbook = load_workbook(uploaded_file, read_only=True, data_only=True)
        return [list(row) for row in workbook.active.iter_rows(values_only=True)]

    raw = uploaded_file.read()
    if isinstance(raw, bytes):
        try:
            text = raw.decode('utf-8-sig')
        except UnicodeDecodeError:
            text = raw.decode('cp1251')
    else:
        text = raw

    sample = text[:2000]
    if '\t' in sample:
        delimiter = '\t'
    elif sample.count(';') >= sample.count(','):
        delimiter = ';'
    else:
        delimiter = ','

    return list(csv.reader(io.StringIO(text), delimiter=delimiter))


def parse_duplicate_rows(rows):
    """Превращает строки файла в список {'serial', 'buyers': [..]}."""
    _, _, parse_date = _helpers()
    items = []
    problems = []

    for line_no, row in enumerate(rows, start=1):
        cells = [_clean(c) for c in row]
        if not cells or not cells[0]:
            continue

        serial = str(cells[0]).upper().replace(' ', '')
        if not SERIAL_RE.match(serial):
            # заголовок или мусор
            continue

        cells = cells + [''] * (10 - len(cells))

        buyers = []
        for start in (1, 6):
            date_raw, model, name, phone = cells[start:start + 4]
            if not any([date_raw, model, name, phone]):
                continue
            sale_date = parse_date(date_raw)
            if date_raw and not sale_date:
                problems.append(f'Строка {line_no}, {serial}: не удалось прочитать дату «{date_raw}»')
            buyers.append({
                'sale_date': sale_date,
                'model_name': str(model) or 'BR Watch',
                'buyer_name': str(name),
                'buyer_phone': str(phone),
            })

        if len(buyers) < 2:
            problems.append(f'Строка {line_no}, {serial}: найден только один покупатель — пропущено')
            continue

        items.append({'serial': serial, 'buyers': buyers})

    return items, problems


def _matches(record, buyer):
    if buyer['buyer_name'] and _norm_name(record.buyer_name) == _norm_name(buyer['buyer_name']):
        return True
    phone = _digits(buyer['buyer_phone'])[-10:]
    if len(phone) >= 10 and _digits(record.buyer_phone)[-10:] == phone:
        return True
    return False


@transaction.atomic
def apply_duplicates(items, dry_run=False):
    """Записывает покупателей в базу. Возвращает статистику."""
    get_auto_color, get_auto_warranty_months, _ = _helpers()
    stats = {'serials': 0, 'created': 0, 'updated': 0}
    log = []

    for item in items:
        serial = item['serial']
        stats['serials'] += 1

        existing = list(
            ProductWarranty.objects.filter(serial_number__iexact=serial).order_by('id')
        )
        used = set()
        plan = []

        # 1) сначала сопоставляем по имени / телефону
        for buyer in item['buyers']:
            match = next((r for r in existing if r.pk not in used and _matches(r, buyer)), None)
            if match:
                used.add(match.pk)
            plan.append([buyer, match])

        # 2) несопоставленным покупателям отдаём свободные записи
        for pair in plan:
            if pair[1] is None:
                free = next((r for r in existing if r.pk not in used), None)
                if free:
                    used.add(free.pk)
                    pair[1] = free

        for buyer, record in plan:
            created = record is None
            if created:
                record = ProductWarranty(serial_number=serial)

            record.buyer_name = buyer['buyer_name']
            record.buyer_phone = buyer['buyer_phone']
            record.model_name = buyer['model_name']
            record.color = get_auto_color(buyer['model_name']) or record.color
            record.warranty_months = get_auto_warranty_months(buyer['model_name'])
            if buyer['sale_date']:
                record.sale_date = buyer['sale_date']
            if record.status == ProductWarranty.Status.IN_STOCK or created:
                record.status = ProductWarranty.Status.SOLD
            if DUPLICATE_NOTE not in (record.notes or ''):
                record.notes = (record.notes + '\n' if record.notes else '') + DUPLICATE_NOTE

            if not dry_run:
                record.save()

            stats['created' if created else 'updated'] += 1
            log.append(
                f'{serial}: {"создана" if created else "обновлена"} запись — '
                f'{buyer["buyer_name"] or "(без имени)"}, {buyer["model_name"]}, '
                f'{buyer["sale_date"].strftime("%d.%m.%Y") if buyer["sale_date"] else "без даты"}'
            )

    if dry_run:
        transaction.set_rollback(True)

    return stats, log
